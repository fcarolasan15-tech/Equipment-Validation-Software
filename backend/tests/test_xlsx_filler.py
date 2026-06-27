import pytest
from pathlib import Path
from openpyxl import Workbook
from app.services.xlsx_filler import fill_xlsx


def test_fill_xlsx_placeholder(tmp_path):
    template = tmp_path / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{equipment_name}}"
    ws["B1"] = "Static"
    wb.save(str(template))

    output = tmp_path / "output.xlsx"
    fill_xlsx(template, output, {"equipment_name": "Retort Machine"})

    from openpyxl import load_workbook
    result = load_workbook(str(output))
    assert result.active["A1"].value == "Retort Machine"
    assert result.active["B1"].value == "Static"
