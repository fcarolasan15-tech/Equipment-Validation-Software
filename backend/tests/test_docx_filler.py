import pytest
from pathlib import Path
from docx import Document
from app.services.docx_filler import fill_docx


@pytest.fixture
def tmp_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Equipment: {{equipment_name}}")
    doc.add_paragraph("Prepared by: {{prepared_by}}")
    doc.add_paragraph("Verified: {{is_verified}}")
    t = doc.add_table(1, 2)
    t.rows[0].cells[0].text = "Status: {{status}}"
    t.rows[0].cells[1].text = "Date: {{date}}"
    path = tmp_path / "test_template.docx"
    doc.save(str(path))
    return path


def test_fill_placeholders(tmp_docx, tmp_path):
    output = tmp_path / "filled.docx"
    data = {
        "equipment_name": "Retort Sterilizer Unit 1",
        "prepared_by": "[Prepared By]",
        "is_verified": True,
        "status": "APPROVED",
        "date": "2025-01-01",
    }
    result = fill_docx(tmp_docx, output, data)
    assert result.exists()

    doc = Document(str(result))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "Retort Sterilizer Unit 1" in full_text
    assert "[Prepared By]" in full_text
    assert "☒" in full_text
    assert "{{equipment_name}}" not in full_text


def test_fill_table_cells(tmp_docx, tmp_path):
    output = tmp_path / "filled_table.docx"
    fill_docx(tmp_docx, output, {"status": "PASS", "date": "2025-06-01"})
    doc = Document(str(output))
    table_text = " ".join(c.text for r in doc.tables[0].rows for c in r.cells)
    assert "PASS" in table_text
    assert "2025-06-01" in table_text


def test_no_placeholder_unchanged(tmp_docx, tmp_path):
    output = tmp_path / "no_change.docx"
    fill_docx(tmp_docx, output, {})
    doc = Document(str(output))
    full = " ".join(p.text for p in doc.paragraphs)
    assert "{{equipment_name}}" in full
