import shutil
from pathlib import Path
from openpyxl import load_workbook


def fill_xlsx(template_path: Path, output_path: Path, data: dict) -> Path:
    shutil.copy2(template_path, output_path)
    wb = load_workbook(str(output_path))

    named_ranges = {}
    for name, ref in wb.defined_names.items():
        try:
            destinations = list(ref.destinations)
            if destinations:
                sheet_name, cell_ref = destinations[0]
                named_ranges[name] = (sheet_name, cell_ref)
        except Exception:
            pass

    for key, value in data.items():
        val = str(value) if value is not None else ""

        if key in named_ranges:
            sheet_name, cell_ref = named_ranges[key]
            clean_ref = cell_ref.replace("$", "")
            if sheet_name in wb.sheetnames:
                wb[sheet_name][clean_ref] = val
            continue

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and f"{{{{{key}}}}}" in cell.value:
                        cell.value = cell.value.replace(f"{{{{{key}}}}}", val)

    wb.save(str(output_path))
    return output_path
