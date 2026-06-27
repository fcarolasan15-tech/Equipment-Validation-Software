import shutil
import re
from pathlib import Path
from docx import Document


CHECKBOX_TRUE = "☒"
CHECKBOX_FALSE = "☐"


def _replace_in_paragraph(para, data: dict):
    full_text = "".join(r.text for r in para.runs)
    if "{{" not in full_text:
        return

    new_text = full_text
    for key, value in data.items():
        placeholder = f"{{{{{key}}}}}"
        if isinstance(value, bool):
            replacement = CHECKBOX_TRUE if value else CHECKBOX_FALSE
        else:
            replacement = str(value) if value is not None else ""
        new_text = new_text.replace(placeholder, replacement)

    if new_text != full_text:
        for i, run in enumerate(para.runs):
            para.runs[i].text = ""
        if para.runs:
            para.runs[0].text = new_text


def fill_docx(template_path: Path, output_path: Path, data: dict) -> Path:
    shutil.copy2(template_path, output_path)
    doc = Document(str(output_path))

    for para in doc.paragraphs:
        _replace_in_paragraph(para, data)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, data)

    for section in doc.sections:
        for para in section.header.paragraphs:
            _replace_in_paragraph(para, data)
        for para in section.footer.paragraphs:
            _replace_in_paragraph(para, data)

    doc.save(str(output_path))
    return output_path
